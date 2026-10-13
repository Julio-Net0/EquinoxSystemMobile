# -*- coding: utf-8 -*-
"""
Gera a apresentação SINTEC 2026 (PowerPoint) do Equinox Mobile.

Usa o "Template-Apresentacao-SINTEC-2026.pptx" como base (mesma estrutura de
slides, logotipos de realizadores/apoiadores e cabeçalho SINTEC) e aplica o
design dos slides HTML do projeto (tema escuro, acentos neon por camada).

Uso:  python scripts/gerar_apresentacao_sintec.py
"""
import os
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(RAIZ, "Template-Apresentacao-SINTEC-2026.pptx")
SAIDA = os.path.join(RAIZ, "Apresentacao-SINTEC-2026-Equinox-Mobile.pptx")

# ---------- Paleta (idêntica aos slides HTML) ----------
BG1, BG2 = "0A1020", "070B14"
PANEL, BORDER = "111827", "1E293B"
TEXT, MUTED, SOFT = "E6EBF5", "94A3B8", "CBD5E1"
CYAN, DOMAIN, APP, ADAPT, UI, TESTS = "7DD3FC", "34D399", "60A5FA", "A78BFA", "FB923C", "F472B6"

FONTE, MONO = "Segoe UI", "Consolas"
TOTAL = 9


def rgb(h):
    return RGBColor.from_string(h)


# ---------- Helpers ----------
def alpha(shape, pct):
    clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct * 1000)))


def back(slide, shape):
    """Envia a forma para o fundo (acima apenas do grupo raiz)."""
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def fundo(slide, glow_baixo=False):
    f = slide.background.fill
    f.gradient()
    f.gradient_angle = 90
    f.gradient_stops[0].color.rgb = rgb(BG1)
    f.gradient_stops[1].color.rgb = rgb(BG2)
    glows = [(-3.0, -3.2, 7.0, CYAN, 14)]
    if glow_baixo:
        glows.append((9.2, 4.6, 6.0, ADAPT, 12))
    for x, y, d, cor, a in reversed(glows):
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        o.fill.solid()
        o.fill.fore_color.rgb = rgb(cor)
        alpha(o, a)
        o.line.fill.background()
        back(slide, o)


def limpar_textos(slide):
    for sh in list(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            sh._element.getparent().remove(sh._element)


def rrect(slide, x, y, w, h, fill=PANEL, line=BORDER, raio=0.08, a=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = raio
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if a is not None:
        alpha(s, a)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def texto(slide, x, y, w, h, partes, tam=18, cor=SOFT, bold=False, fonte=FONTE,
          alinh=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, espaco=1.1, entre=0):
    """partes: str | lista de parágrafos; cada parágrafo = str | lista de (txt, cor, bold)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(partes, str):
        partes = [partes]
    for i, par in enumerate(partes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinh
        p.line_spacing = espaco
        if i:
            p.space_before = Pt(entre)
        runs = [(par, cor, bold)] if isinstance(par, str) else par
        for t, c, b in runs:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(tam)
            r.font.name = fonte
            r.font.bold = b
            r.font.color.rgb = rgb(c)
    return tb


def cabecalho(slide, rotulo, idx, acento, titulo_runs, y_titulo=0.95, tam=38):
    # selo (badge) com ponto luminoso
    larg = 0.55 + 0.135 * len(rotulo)
    rrect(slide, 0.6, 0.38, larg, 0.4, fill=acento, line=acento, raio=0.5, a=12)
    d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.78), Inches(0.53), Inches(0.1), Inches(0.1))
    d.fill.solid(); d.fill.fore_color.rgb = rgb(acento); d.line.fill.background()
    texto(slide, 1.0, 0.38, larg - 0.4, 0.4, rotulo.upper(), tam=12, cor=acento, bold=True,
          anchor=MSO_ANCHOR.MIDDLE)
    # numeração mono
    texto(slide, 10.9, 0.38, 1.85, 0.4,
          [[(f"{idx:02d}", acento, True), (f" / {TOTAL:02d}", MUTED, False)]],
          tam=14, fonte=MONO, alinh=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    texto(slide, 0.6, y_titulo, 12.1, 0.8, [titulo_runs], tam=tam, bold=True, cor="FFFFFF")


def rodape(slide, idx, brand=True, barra=True):
    # barra de progresso multicolorida (mesma do HUD dos slides HTML), agora no TOPO:
    # o rodapé pertence à faixa de logos do template.
    cores = [DOMAIN, APP, ADAPT, UI, TESTS]
    total_w = 13.333
    rect(slide, 0, 0, total_w, 0.07, "1B2435")
    fr = idx / TOTAL
    seg = total_w * fr / len(cores)
    for i, c in enumerate(cores):
        rect(slide, seg * i, 0, seg + 0.01, 0.07, c)


LIMITE_CONTEUDO = 5.95  # polegadas: a faixa de logos começa em 6.1


def reescalar(slide, limiar, origem_antiga, origem_nova, kg, kf):
    """Comprime verticalmente o conteúdo (formas com topo >= limiar) para caber acima dos logos."""
    for sh in slide.shapes:
        t = sh.top / 914400
        if t < limiar:
            continue
        sh.top = Inches(origem_nova + (t - origem_antiga) * kg)
        sh.height = Inches(sh.height / 914400 * kg)
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        r.font.size = Pt(max(10, round(r.font.size.pt * kf * 2) / 2))


def faixa_logos_no_layout(prs, slide_ref):
    """Insere uma faixa branca no layout BLANK, logo abaixo dos logos/rótulos do template."""
    layout = slide_ref.slide_layout
    primeiro = [sh for sh in layout.shapes if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM][0]
    tmp = rect(slide_ref, 0, 6.1, 13.333, 1.4, "FFFFFF")
    el = tmp._element
    el.getparent().remove(el)
    el.find(".//" + qn("p:cNvPr")).set("id", "9999")
    primeiro._element.addprevious(el)


def card(slide, x, y, w, h, acento, titulo, corpo, tam_t=17, tam_c=14, num=None):
    rrect(slide, x, y, w, h)
    rect(slide, x, y + 0.25, 0.06, h - 0.5, acento)
    ty = y + 0.2
    if num:
        texto(slide, x + 0.3, ty, 1, 0.3, num, tam=13, cor=acento, bold=True, fonte=MONO)
        ty += 0.34
    texto(slide, x + 0.3, ty, w - 0.55, 0.5, titulo, tam=tam_t, cor="FFFFFF", bold=True)
    texto(slide, x + 0.3, ty + 0.46, w - 0.55, h - (ty - y) - 0.6, corpo, tam=tam_c, cor=SOFT, espaco=1.15)


def bullets(slide, x, y, w, itens, acento, tam=18, passo=0.85):
    for i, it in enumerate(itens):
        yy = y + i * passo
        b = rect(slide, x, yy + 0.16, 0.2, 0.045, acento)
        texto(slide, x + 0.4, yy, w - 0.4, passo, [it] if isinstance(it, list) else it,
              tam=tam, cor=SOFT, espaco=1.12)


def chip(slide, x, y, w, h, txt, cor, tam=14):
    rrect(slide, x, y, w, h, fill="0F172A", line=cor, raio=0.3)
    texto(slide, x, y, w, h, txt, tam=tam, cor=cor, bold=True, alinh=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)


def seta(slide, x, y, w=0.4, cor=MUTED):
    texto(slide, x, y, w, 0.4, "→", tam=20, cor=cor, alinh=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def novo_slide(prs, layout):
    return prs.slides.add_slide(layout)


# ================== Construção ==================
def main():
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    blank = S[0].slide_layout

    # ---------- 1. CAPA ----------
    s = S[0]
    foto = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    limpar_textos(s)
    fundo(s)
    cartao = rrect(s, 0.6, 0.45, 4.35, 1.85, fill="FFFFFF", line=None, raio=0.1)
    foto._element.addprevious(cartao._element)
    foto.left, foto.top = Inches(0.78), Inches(0.62)
    foto.width, foto.height = Inches(4.0), Inches(1.51)
    texto(s, 5.35, 0.62, 7.4, 1.5,
          [[("SIMPÓSIO DE", CYAN, True)], [("INOVAÇÃO, TECNOLOGIA, EMPREENDEDORISMO", TEXT, True)],
           [("E CIÊNCIAS COMPUTACIONAIS", TEXT, True)]],
          tam=16, anchor=MSO_ANCHOR.MIDDLE, espaco=1.2)
    texto(s, 5.35, 1.95, 7.4, 0.35, "SINTEC 2026", tam=14, cor=MUTED, fonte=MONO)

    texto(s, 0.6, 2.75, 12.1, 2.0,
          [[("EQUINOX MOBILE: ", CYAN, True),
            ("Aplicativo Móvel Multi-tenant e Offline-First para Gestão Georreferenciada "
             "de Usinas Fotovoltaicas", "FFFFFF", True)]],
          tam=33, espaco=1.05)
    rrect(s, 0.6, 4.7, 7.3, 0.46, fill=ADAPT, line=ADAPT, raio=0.5, a=14)
    texto(s, 0.6, 4.7, 7.3, 0.46,
          [[("ÁREA TEMÁTICA  ", ADAPT, True), ("Projeto de Extensão", TEXT, False)]],
          tam=14, alinh=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    autores = [("Victor Couto Zanaroli", "Autor e Apresentador", "victorzanaroli@gmail.com", DOMAIN),
               ("Júlio Ignácio Miranda Neto", "Autor e Apresentador", "jjulio2468@gmail.com", APP),
               ("Lázaro Eduardo da Silva", "Orientador", "lazaro@cefetmg.br", UI)]
    for i, (n, papel, mail, c) in enumerate(autores):
        x = 0.6 + i * 4.1
        rrect(s, x, 5.5, 3.9, 1.3)
        rect(s, x, 5.7, 0.06, 0.9, c)
        texto(s, x + 0.3, 5.62, 3.5, 0.3, papel.upper(), tam=11, cor=c, bold=True, fonte=MONO)
        texto(s, x + 0.3, 5.92, 3.5, 0.4, n, tam=17, cor="FFFFFF", bold=True)
        texto(s, x + 0.3, 6.38, 3.5, 0.3, mail, tam=12, cor=MUTED, fonte=MONO)
    rodape(s, 1, brand=False)

    # ---------- 2. OBJETIVO ----------
    s = S[1]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Objetivo", 2, CYAN, [("Digitalizar a coleta, ", "FFFFFF", True),
                                        ("sem depender de internet", CYAN, True)], tam=34)
    rrect(s, 0.6, 1.95, 5.0, 4.85)
    rect(s, 0.6, 2.25, 0.06, 4.25, CYAN)
    texto(s, 0.95, 2.2, 4.4, 0.3, "OBJETIVO GERAL", tam=12, cor=CYAN, bold=True, fonte=MONO)
    texto(s, 0.95, 2.65, 4.4, 4.0,
          "Desenvolver um aplicativo móvel multi-tenant e offline-first que digitalize a coleta de "
          "consumo e geração fotovoltaica (kWh) em usinas de regiões rurais e remotas, garantindo "
          "integridade dos dados e auditoria georreferenciada sem conectividade contínua.",
          tam=19, cor=TEXT, espaco=1.2)
    objs = [("01", "Operação offline", "Persistência local em SQLite com UUIDv4 gerados no domínio, sem colisão de dados.", DOMAIN),
            ("02", "Comprovação de presença", "Câmera e GPS nativos, com compressão automática das fotos dos medidores.", APP),
            ("03", "Sincronização resiliente", "Fila operacional (action queue) enviando à nuvem Supabase quando a rede volta.", ADAPT),
            ("04", "Gestão multi-tenant", "Empresas, usinas e perfis de usuário com autorização controlada.", UI)]
    for i, (n, t, c, cor) in enumerate(objs):
        x = 5.9 + (i % 2) * 3.52
        y = 1.95 + (i // 2) * 2.5
        card(s, x, y, 3.32, 2.35, cor, t, c, tam_t=14, tam_c=13, num=n)
    rodape(s, 2)

    # ---------- 3. JUSTIFICATIVA ----------
    s = S[2]
    limpar_textos(s)
    for sh in list(s.shapes):  # logos já vêm do layout; evita duplicados
        if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM:
            sh._element.getparent().remove(sh._element)
    fundo(s)
    cabecalho(s, "Justificativa", 3, UI, [("Onde a rede falha, ", "FFFFFF", True),
                                          ("o dado não pode falhar", UI, True)], tam=34)
    itens = [
        [("Conectividade instável. ", "FFFFFF", True), ("Usinas em regiões rurais e remotas operam com rede ausente ou intermitente.", SOFT, False)],
        [("Faturamento depende da leitura. ", "FFFFFF", True), ("O consumo e a geração em kWh precisam ser registrados com confiabilidade.", SOFT, False)],
        [("Auditoria no campo. ", "FFFFFF", True), ("Foto do medidor e coordenadas GPS comprovam a presença e a medição.", SOFT, False)],
        [("Múltiplas empresas. ", "FFFFFF", True), ("Um único app atende vários clientes, isolando dados e perfis (multi-tenant).", SOFT, False)],
    ]
    bullets(s, 0.6, 1.95, 7.4, itens, UI, tam=16, passo=0.95)
    rrect(s, 8.5, 2.0, 4.23, 3.75)
    rect(s, 8.5, 2.3, 0.06, 3.15, UI)
    texto(s, 8.85, 2.25, 3.7, 0.3, "A PROPOSTA", tam=12, cor=UI, bold=True, fonte=MONO)
    texto(s, 8.85, 2.7, 3.7, 2.9,
          [[("Offline-first", UI, True), (": o dispositivo é a fonte única de verdade em campo; a nuvem "
            "recebe os dados depois, de forma assíncrona e sem perdas.", TEXT, False)]],
          tam=18, espaco=1.2)
    rodape(s, 3)

    # ---------- 4. METODOLOGIA ----------
    s = S[3]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Metodologia", 4, APP, [("Engenharia de software ", "FFFFFF", True),
                                          ("rigorosa, de ponta a ponta", APP, True)], tam=34)
    fluxo = [("Clean Architecture", DOMAIN), ("DDD", APP), ("TDD", TESTS), ("Expo + Supabase", UI)]
    x = 0.6
    for i, (t, c) in enumerate(fluxo):
        w = 0.5 + 0.13 * len(t)
        chip(s, x, 1.95, w, 0.48, t, c, tam=14)
        x += w
        if i < len(fluxo) - 1:
            seta(s, x, 1.99, 0.4)
            x += 0.4
    cols = [("Arquitetura Limpa", "Camadas isoladas: domínio, aplicação, infraestrutura e apresentação. "
             "A regra de negócio não conhece banco, rede nem interface.", DOMAIN),
            ("Domain-Driven Design", "Entidades Usina, Leitura, Empresa e Usuário; value objects CNPJ, "
             "CoordenadasGPS, UUIDv4 e ValorKwh.", APP),
            ("Test-Driven Development", "Testes escritos antes do código com Jest e Testing Library, "
             "cobrindo domínio, casos de uso, infraestrutura e telas.", TESTS),
            ("Stack tecnológica", "React Native com Expo SDK 57 e TypeScript estrito; SQLite local; "
             "Supabase (Auth, Storage, PostgreSQL); Zod e React Hook Form.", UI)]
    for i, (t, c, cor) in enumerate(cols):
        card(s, 0.6 + i * 3.07, 2.75, 2.85, 3.95, cor, t, c, tam_t=14, tam_c=14)
    texto(s, 0.6, 6.78, 12.1, 0.3,
          "A inteligência artificial apoiou a estruturação arquitetural e o refinamento da redação técnica.",
          tam=12, cor=MUTED)
    rodape(s, 4)

    # ---------- 5. DESCRIÇÃO (1/2) — Arquitetura ----------
    s = S[4]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Descrição do projeto · 1/2", 5, ADAPT,
              [("Arquitetura em ", "FFFFFF", True), ("camadas", ADAPT, True)], tam=36)
    camadas = [("APRESENTAÇÃO", UI, "Telas Expo Router, componentes e hooks controladores (12 telas)"),
               ("APLICAÇÃO", APP, "Casos de uso: registrar leitura, cadastrar usina, autorizar usuário… (8)"),
               ("DOMÍNIO", DOMAIN, "Entidades, value objects, serviços e portas (interfaces) puras"),
               ("INFRAESTRUTURA", ADAPT, "SQLite, Supabase, câmera, GPS, SecureStore e monitor de rede")]
    for i, (n, c, d) in enumerate(camadas):
        y = 2.0 + i * 1.18
        rrect(s, 0.6, y, 7.5, 1.02)
        rrect(s, 0.6, y, 2.35, 1.02, fill=c, line=c, raio=0.1)
        texto(s, 0.6, y, 2.35, 1.02, n, tam=14, cor="0B1020", bold=True, alinh=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
        texto(s, 3.2, y, 4.7, 1.02, d, tam=15, cor=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    rrect(s, 8.5, 2.0, 4.23, 4.6)
    rect(s, 8.5, 2.3, 0.06, 4.0, ADAPT)
    texto(s, 8.85, 2.25, 3.7, 0.3, "REGRA DE DEPENDÊNCIA", tam=12, cor=ADAPT, bold=True, fonte=MONO)
    texto(s, 8.85, 2.7, 3.7, 2.0,
          [[("Tudo aponta para o ", TEXT, False), ("domínio", DOMAIN, True), (".", TEXT, False)]],
          tam=22, bold=True, espaco=1.15)
    texto(s, 8.85, 3.65, 3.7, 2.9,
          "Infraestrutura implementa as portas definidas no domínio, o que permite trocar "
          "SQLite, Supabase ou hardware sem alterar regras de negócio e testar tudo com fakes em memória.",
          tam=15, cor=SOFT, espaco=1.2)
    rodape(s, 5)

    # ---------- 6. DESCRIÇÃO (2/2) — Offline-first (novo slide) ----------
    s2 = novo_slide(prs, blank)
    fundo(s2)
    cabecalho(s2, "Descrição do projeto · 2/2", 6, ADAPT,
              [("Motor ", "FFFFFF", True), ("offline-first", ADAPT, True), (" e sincronização", "FFFFFF", True)],
              tam=34)
    nos = [("Leitura", "foto + GPS", DOMAIN), ("SQLite", "UUIDv4 local", APP),
           ("Action Queue", "fila de ações", ADAPT), ("SyncManager", "rede restabelecida", UI),
           ("Supabase", "nuvem", TESTS)]
    w, gap = 2.0, 0.53
    for i, (t, sub, c) in enumerate(nos):
        x = 0.6 + i * (w + gap)
        rrect(s2, x, 2.0, w, 1.3, fill="0F172A", line=c, raio=0.1)
        texto(s2, x, 2.2, w, 0.4, t, tam=17, cor=c, bold=True, alinh=PP_ALIGN.CENTER)
        texto(s2, x, 2.7, w, 0.4, sub, tam=12, cor=MUTED, alinh=PP_ALIGN.CENTER, fonte=MONO)
        if i < len(nos) - 1:
            seta(s2, x + w + 0.03, 2.45, gap - 0.06, cor=MUTED)
    trio = [("Compressão de mídia", "Fotos dos medidores são comprimidas automaticamente em segundo plano, "
             "reduzindo o espaço local e o tráfego de envio.", DOMAIN),
            ("Resolução de conflitos", "Sincronização incremental com resolução automática: prevalece o "
             "registro com data e hora mais recente.", APP),
            ("Perfis e autorização", "Cadastro de usuários pendentes, autorização por empresa e gestão de "
             "estrutura (usinas e equipes) pelo painel administrativo.", UI)]
    for i, (t, c, cor) in enumerate(trio):
        card(s2, 0.6 + i * 4.1, 3.75, 3.9, 2.95, cor, t, c, tam_t=17, tam_c=14)
    rodape(s2, 6)

    # ---------- 7. FOTOS / PRINTS ----------
    s = S[5]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Andamento do projeto", 7, TESTS,
              [("Prints e ", "FFFFFF", True), ("evidências do andamento", TESTS, True)], tam=34)
    stats = [("12", "telas"), ("8", "casos de uso"), ("5", "entidades de domínio"), ("30", "arquivos de teste")]
    for i, (n, l) in enumerate(stats):
        x = 0.6 + i * 3.07
        rrect(s, x, 1.95, 2.85, 1.15)
        texto(s, x + 0.3, 2.0, 1.3, 1.05, n, tam=38, cor=CYAN, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        texto(s, x + 1.35, 2.0, 1.4, 1.05, l.upper(), tam=11, cor=MUTED, fonte=MONO,
              anchor=MSO_ANCHOR.MIDDLE)
    telas = ["Login", "Painel / Dashboard", "Nova leitura (câmera + GPS)", "Administração"]
    cores = [DOMAIN, APP, UI, ADAPT]
    for i, (t, c) in enumerate(zip(telas, cores)):
        x = 0.6 + i * 3.07
        fr = rrect(s, x + 0.45, 3.3, 1.95, 3.15, fill="0F172A", line=c, raio=0.12)
        rrect(s, x + 1.05, 3.4, 0.75, 0.07, fill=BORDER, line=None, raio=0.5)
        texto(s, x + 0.45, 3.3, 1.95, 3.15, [[("INSERIR PRINT", c, True)], [("da tela", MUTED, False)]],
              tam=12, fonte=MONO, alinh=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        texto(s, x, 6.55, 2.85, 0.4, t, tam=13, cor=SOFT, bold=True, alinh=PP_ALIGN.CENTER)
    rodape(s, 7)

    # ---------- 8. CONSIDERAÇÕES FINAIS ----------
    s = S[6]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Considerações finais", 8, DOMAIN,
              [("Confiabilidade ", "FFFFFF", True), ("operacional no campo", DOMAIN, True)], tam=34)
    res = [("Fluxo de trabalho resiliente", "Leituras offline executadas de forma fluida, sem travar a operação.", DOMAIN),
           ("Mídias otimizadas", "Redução expressiva do tamanho das fotos armazenadas.", APP),
           ("Sincronização eficiente", "Envio completo e automático após o restabelecimento da rede.", ADAPT)]
    for i, (t, c, cor) in enumerate(res):
        card(s, 0.6, 1.95 + i * 1.65, 6.4, 1.5, cor, t, c, tam_t=17, tam_c=14)
    rrect(s, 7.35, 1.95, 5.38, 4.8)
    rect(s, 7.35, 2.25, 0.06, 4.2, DOMAIN)
    texto(s, 7.7, 2.2, 4.8, 0.3, "CONCLUSÃO", tam=12, cor=DOMAIN, bold=True, fonte=MONO)
    texto(s, 7.7, 2.65, 4.8, 4.0,
          [[("A abordagem ", TEXT, False), ("offline-first", DOMAIN, True),
            (" combinada a padrões rigorosos de engenharia de software ", TEXT, False),
            ("elimina falhas na coleta de medições", DOMAIN, True),
            (" no campo, oferecendo alta confiabilidade operacional para a gestão distribuída "
             "de energia solar.", TEXT, False)]],
          tam=19, espaco=1.22)
    rodape(s, 8)

    # ---------- 9. REFERÊNCIAS ----------
    s = S[7]
    limpar_textos(s); fundo(s)
    cabecalho(s, "Referências bibliográficas", 9, CYAN,
              [("Referências ", "FFFFFF", True), ("bibliográficas", CYAN, True)], tam=34)
    refs = [
        ("BECK, Kent.", " Test-Driven Development: By Example. Boston: Addison-Wesley, 2002."),
        ("EVANS, Eric.", " Domain-Driven Design: Tackling Complexity in the Heart of Software. Boston: Addison-Wesley, 2003."),
        ("EXPO.", " Expo Documentation - SDK 57. 2026. Disponível em: docs.expo.dev/versions/v57.0.0. Acesso em: 17 set. 2026."),
        ("FOWLER, Martin.", " Patterns of Enterprise Application Architecture. Boston: Addison-Wesley, 2002."),
        ("MARTIN, Robert C.", " Clean Architecture: A Craftsman's Guide to Software Structure and Design. Upper Saddle River: Prentice Hall, 2017."),
        ("REACT NATIVE.", " React Native: Learn the basics. 2026. Disponível em: reactnative.dev/docs/getting-started. Acesso em: 17 set. 2026."),
        ("SQLITE.", " SQLite Documentation. 2026. Disponível em: sqlite.org/docs.html. Acesso em: 17 set. 2026."),
        ("SUPABASE.", " Supabase Documentation: Database & Authentication. 2026. Disponível em: supabase.com/docs. Acesso em: 17 set. 2026."),
    ]
    acentos = [DOMAIN, APP, ADAPT, UI]
    for i, (a, b) in enumerate(refs):
        col, lin = i // 4, i % 4
        x, y = 0.6 + col * 6.15, 1.95 + lin * 1.2
        rrect(s, x, y, 5.95, 1.05)
        rect(s, x, y + 0.2, 0.06, 0.65, acentos[lin])
        texto(s, x + 0.3, y, 5.45, 1.05, [[(f"{i + 1}. ", acentos[lin], True), (a, "FFFFFF", True), (b, SOFT, False)]],
              tam=12, anchor=MSO_ANCHOR.MIDDLE, espaco=1.1)
    rodape(s, 9)

    # Comprime o conteúdo para ficar acima da faixa de logos (que vem do layout)
    for sl in (S[1], S[3], S[4], S[5], S[6], S[7], s2):
        reescalar(sl, 1.9, 1.95, 1.8, 0.83, 0.95)
    reescalar(S[0], 2.6, 2.75, 2.5, 0.85, 0.93)
    faixa_logos_no_layout(prs, S[2])

    # Move o slide novo (índice 8) para logo após a descrição 1/2 (posição 6)
    lst = prs.slides._sldIdLst
    el = lst[8]
    lst.remove(el)
    lst.insert(5, el)

    prs.save(SAIDA)
    print("OK ->", SAIDA)


if __name__ == "__main__":
    main()
